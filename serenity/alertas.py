import os
import smtplib
from datetime import datetime
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import pandas as pd
import psycopg2
from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart, Reference

from .config import (
    GMAIL_PASS,
    GMAIL_USER,
    TWILIO_WHATSAPP_FROM,
    TWILIO_WHATSAPP_TO,
    logger,
    twilio_client,
)


def _db_conn():
    database_url = os.getenv("DATABASE_URL", "").strip()
    if not database_url:
        raise RuntimeError("DATABASE_URL no configurada.")
    return psycopg2.connect(database_url)


def _obtener_alerta_y_contexto(user_id):
    conn = _db_conn()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT a.id, a.tema, a.nivel_riesgo, a.razon,
                       TO_CHAR(a.timestamp AT TIME ZONE 'America/Mexico_City', 'DD/MM/YYYY HH24:MI') AS fecha,
                       COALESCE(ac.mensaje_detonante, '') AS detonante,
                       ac.mensajes_previos_json
                FROM alertas a
                LEFT JOIN alertas_contexto ac ON ac.alerta_id = a.id
                WHERE a.user_id=%s
                ORDER BY a.timestamp DESC
                LIMIT 1
                """,
                (user_id,),
            )
            row = cur.fetchone()
            if not row:
                return {
                    "tema": "sin_tema",
                    "nivel": "medio",
                    "razon": "",
                    "fecha": "",
                    "detonante": "",
                    "previos": [],
                }
            previos = row[6] or []
            if isinstance(previos, str):
                try:
                    import json
                    previos = json.loads(previos)
                except Exception:
                    previos = []
            return {
                "alerta_id": row[0],
                "tema": row[1] or "sin_tema",
                "nivel": row[2] or "medio",
                "razon": row[3] or "",
                "fecha": row[4] or "",
                "detonante": row[5] or "",
                "previos": previos,
            }
    finally:
        conn.close()


def _render_mensaje_estilo_gmail(usuario_info, alerta):
    nombre = usuario_info.get("nombre_completo") or usuario_info.get("username") or str(usuario_info.get("id", "N/A"))
    tema = str(alerta.get("tema", "sin_tema")).capitalize()
    nivel = str(alerta.get("nivel", "medio")).upper()
    detonante = alerta.get("detonante") or alerta.get("razon") or "Sin mensaje detonante registrado."
    previos = alerta.get("previos", [])[:5]

    lineas_previas = []
    if previos:
        for i, m in enumerate(previos, start=1):
            fecha = m.get("fecha", "--:--")
            texto = (m.get("mensaje", "") or "").strip()
            lineas_previas.append(f"{i}.- [{fecha}] {texto}")
    lineas_previas.append(f"{len(lineas_previas)+1}.- [detonante] {detonante}")

    texto_plano = (
        f"🚨 ALERTA DE {nivel} RIESGO: {tema.upper()} 🚨\n\n"
        f"👤 Usuario: {nombre}\n"
        f"🧾 Matrícula: {usuario_info.get('matricula', '-')}\n"
        f"🏫 Facultad: {usuario_info.get('facultad', '-')}\n"
        f"📧 Correo: {usuario_info.get('correo', '-')}\n"
        f"📱 Teléfono: {usuario_info.get('telefono', '-')}\n"
        f"🕒 Fecha de alerta: {alerta.get('fecha', '-')}\n\n"
        f"🚨 Mensaje que activó la alerta:\n\"{detonante}\"\n\n"
        f"📅 Últimos 5 mensajes recientes:\n" + "\n".join(lineas_previas)
    )

    items_html = "".join([f"<li>{linea}</li>" for linea in lineas_previas])
    html = f"""
    <html>
      <body style="font-family:Arial,Helvetica,sans-serif;color:#222;line-height:1.45;">
        <h2 style="margin:0 0 10px 0;">Alerta de riesgo - Usuario {nombre} - Tema: {tema}</h2>
        <p style="font-size:22px;margin:12px 0;">🚨 ALERTA DE {nivel} RIESGO: {tema.upper()} 🚨</p>
        <p><b>👤 Usuario:</b> {nombre}</p>
        <p><b>🧾 Matrícula:</b> {usuario_info.get('matricula', '-')}</p>
        <p><b>🏫 Facultad:</b> {usuario_info.get('facultad', '-')}</p>
        <p><b>📧 Correo:</b> {usuario_info.get('correo', '-')}</p>
        <p><b>📱 Teléfono:</b> {usuario_info.get('telefono', '-')}</p>
        <p><b>🕒 Fecha de alerta:</b> {alerta.get('fecha', '-')}</p>
        <p style="margin-top:16px;"><b>🚨 Mensaje que activó la alerta:</b><br>"{detonante}"</p>
        <p style="margin-top:16px;"><b>📅 Últimos 5 mensajes recientes:</b></p>
        <ol style="margin-top:4px;">{items_html}</ol>
      </body>
    </html>
    """
    return texto_plano, html


def generar_excel_clasificacion(usuario_info):
    conn = _db_conn()
    try:
        df_tema = pd.read_sql_query(
            """
            SELECT tema, COUNT(*) AS total
            FROM alertas
            WHERE user_id = %s
            GROUP BY tema
            ORDER BY total DESC
            """,
            conn,
            params=(usuario_info["id"],),
        )
        df_nivel = pd.read_sql_query(
            """
            SELECT nivel_riesgo, COUNT(*) AS total
            FROM alertas
            WHERE user_id = %s
            GROUP BY nivel_riesgo
            ORDER BY total DESC
            """,
            conn,
            params=(usuario_info["id"],),
        )

        if df_tema.empty:
            df_tema = pd.DataFrame([{"tema": "sin_datos", "total": 0}])
        if df_nivel.empty:
            df_nivel = pd.DataFrame([{"nivel_riesgo": "sin_datos", "total": 0}])

        archivo = f"clasificacion_serenity_{usuario_info['id']}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        with pd.ExcelWriter(archivo, engine="openpyxl") as writer:
            df_tema.to_excel(writer, sheet_name="clasif_tema", index=False)
            df_nivel.to_excel(writer, sheet_name="clasif_nivel", index=False)

        wb = load_workbook(archivo)
        ws_tema = wb["clasif_tema"]
        ws_nivel = wb["clasif_nivel"]

        bar = BarChart()
        bar.title = "Interacciones por tema"
        bar.y_axis.title = "Total"
        bar.x_axis.title = "Tema"
        bar.add_data(Reference(ws_tema, min_col=2, min_row=1, max_row=ws_tema.max_row), titles_from_data=True)
        bar.set_categories(Reference(ws_tema, min_col=1, min_row=2, max_row=ws_tema.max_row))
        ws_tema.add_chart(bar, "E2")

        pie = PieChart()
        pie.title = "Distribucion por nivel de riesgo"
        pie.add_data(Reference(ws_nivel, min_col=2, min_row=1, max_row=ws_nivel.max_row), titles_from_data=True)
        pie.set_categories(Reference(ws_nivel, min_col=1, min_row=2, max_row=ws_nivel.max_row))
        ws_nivel.add_chart(pie, "E2")
        wb.save(archivo)
        return archivo
    finally:
        conn.close()


def enviar_correo_alerta(destinatario, usuario_info, motivo_riesgo):
    if not GMAIL_USER or not GMAIL_PASS:
        raise RuntimeError("No se han configurado las credenciales de Gmail.")

    alerta = _obtener_alerta_y_contexto(usuario_info["id"])
    texto_plano, html = _render_mensaje_estilo_gmail(usuario_info, alerta)

    archivo_excel = f"alerta_{usuario_info['id']}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    archivo_clasificacion = generar_excel_clasificacion(usuario_info)
    pd.DataFrame([usuario_info]).to_excel(archivo_excel, index=False)

    nombre = usuario_info.get("nombre_completo") or usuario_info.get("username") or str(usuario_info.get("id", "N/A"))
    tema = str(alerta.get("tema", "sin_tema")).capitalize()

    msg = MIMEMultipart("alternative")
    msg["From"] = GMAIL_USER
    msg["To"] = destinatario
    msg["Subject"] = f"Alerta de riesgo - Usuario {nombre} - Tema: {tema}"
    msg.attach(MIMEText(texto_plano, "plain", "utf-8"))
    msg.attach(MIMEText(html, "html", "utf-8"))

    mixed = MIMEMultipart()
    mixed["From"] = msg["From"]
    mixed["To"] = msg["To"]
    mixed["Subject"] = msg["Subject"]
    mixed.attach(msg)

    try:
        for path in (archivo_excel, archivo_clasificacion):
            with open(path, "rb") as fh:
                part = MIMEBase("application", "octet-stream")
                part.set_payload(fh.read())
            encoders.encode_base64(part)
            part.add_header("Content-Disposition", f"attachment; filename={os.path.basename(path)}")
            mixed.attach(part)

        server = smtplib.SMTP("smtp.gmail.com", 587)
        server.starttls()
        server.login(GMAIL_USER, GMAIL_PASS)
        server.send_message(mixed)
        server.quit()
        logger.info(f"Correo de alerta enviado a {destinatario}")
    except Exception as e:
        logger.error(f"Error enviando correo: {e}")
        raise
    finally:
        for path in (archivo_excel, archivo_clasificacion):
            try:
                if os.path.exists(path):
                    os.remove(path)
            except Exception:
                pass


def enviar_whatsapp_emergencia(mensaje):
    if not twilio_client:
        logger.error("Twilio no esta configurado.")
        return
    try:
        message = twilio_client.messages.create(
            from_=TWILIO_WHATSAPP_FROM,
            body=f"SERENITY ALERT: {mensaje}",
            to=TWILIO_WHATSAPP_TO,
        )
        logger.info(f"WhatsApp enviado con exito. SID: {message.sid}")
    except Exception as e:
        logger.error(f"Error enviando WhatsApp: {e}")
